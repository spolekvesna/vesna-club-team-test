from __future__ import annotations
import os, re, json
from pathlib import Path
from datetime import date
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select, func
from openpyxl import Workbook

# DATABASE_URL is set by the test command before import.
from app.main import app, sync_indicator_records
from app.db import SessionLocal, Base, engine
from app.models import *
from app.security import hash_password
from app.indicators import indicator_clients, supported_clients
from app.backups import perform_backup, test_restore as verify_restore
from app.importer import save_import_upload, apply_import

@pytest.fixture(scope='module', autouse=True)
def reset_db():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    with TestClient(app) as c:
        pass
    yield


def login(client: TestClient, email='admin@example.test'):
    r=client.get('/login'); assert r.status_code==200
    csrf=re.search(r'name="csrf" value="([^"]+)"',r.text).group(1)
    r=client.post('/login',data={'csrf':csrf,'email':email,'password':'VesnaDemo!2026'},follow_redirects=False)
    assert r.status_code==303


def csrf(client: TestClient, path='/'):
    r=client.get(path); assert r.status_code==200
    m=re.search(r'name="csrf" value="([^"]+)"',r.text)
    return m.group(1) if m else client.cookies.get('csrf','')


def make_client(db, code='T-001', name='Test'):
    c=Client(client_id='CLI-'+code,code=code,first_name=name,last_name='Klient',birth_date=date(1990,1,1),entry_date=date(2026,10,1),status='active')
    db.add(c); db.commit(); db.refresh(c); return c


def get_worker(db,email='admin@example.test'):
    return db.scalar(select(Worker).where(Worker.email==email))


def test_catalog_seed_216():
    with SessionLocal() as db:
        assert db.scalar(select(func.count()).select_from(CatalogItem)) == 216
        assert db.scalar(select(func.count()).select_from(CatalogItem).where(CatalogItem.category_key=='Oblast')) == 14
        assert db.scalar(select(func.count()).select_from(CatalogItem).where(CatalogItem.category_key=='Cinnost')) == 118


def test_indicator_unique_person_across_multiple_ka():
    with SessionLocal() as db:
        w=get_worker(db); c=make_client(db,'T-IND')
        s=IndividualService(service_id='S-IND',client_id=c.id,service_date=date(2026,10,2),minutes=30,worker_id=w.id,ka='KA3',subactivity='3.1',eligible_project=True)
        db.add(s); db.flush()
        e=Event(event_id='E-IND',title='Skupina',event_date=date(2026,10,3),ka='KA2',subactivity='2.1',event_type='group_program',created_by_id=w.id)
        db.add(e); db.flush(); db.add(EventParticipant(event_id=e.id,client_id=c.id,participant_kind='project',attendance_status='present',eligible_project=True)); db.commit()
        ids=supported_clients(db)
        assert ids.count(c.id)==1


def test_specific_indicator_not_inferred_from_general_record():
    with SessionLocal() as db:
        w=get_worker(db); c=make_client(db,'T-SPEC')
        s=IndividualService(service_id='S-GENERAL',client_id=c.id,service_date=date(2026,10,4),minutes=20,worker_id=w.id,ka='KA1',area_id='OBL001',eligible_project=True)
        db.add(s); db.commit()
        assert c.id not in indicator_clients(db,'026200')
        sync_indicator_records(db,'service',s.service_id,s.service_date,[c.id],['026200'],True,'PL-X'); db.commit()
        assert c.id in indicator_clients(db,'026200')


def test_inactive_catalog_keeps_history():
    with SessionLocal() as db:
        w=get_worker(db); c=make_client(db,'T-CAT')
        item=db.scalar(select(CatalogItem).where(CatalogItem.item_id=='CIN001')); assert item
        s=IndividualService(service_id='S-CAT',client_id=c.id,service_date=date(2026,10,5),minutes=15,worker_id=w.id,activity_id=item.item_id,eligible_project=True)
        db.add(s); db.commit(); item.active=False; item.version+=1; db.commit(); db.refresh(s)
        assert s.activity_id=='CIN001'
        assert db.scalar(select(CatalogItem.name).where(CatalogItem.item_id==s.activity_id))


def test_concurrent_edit_returns_409():
    with SessionLocal() as db:
        c=make_client(db,'T-CONC')
        cid=c.id
    a=TestClient(app); b=TestClient(app); login(a); login(b)
    ra=a.get(f'/clients/{cid}/edit'); rb=b.get(f'/clients/{cid}/edit')
    ca=re.search(r'name="csrf" value="([^"]+)"',ra.text).group(1); cb=re.search(r'name="csrf" value="([^"]+)"',rb.text).group(1)
    va=re.search(r'name="version" value="(\d+)"',ra.text).group(1); vb=re.search(r'name="version" value="(\d+)"',rb.text).group(1)
    payload={'code':'T-CONC','first_name':'První','last_name':'Klient','birth_date':'1990-01-01','entry_date':'2026-10-01','contact':'','fewer':'','document_status':'unknown','drive_ref':'','status':'active','notes':''}
    r=a.post(f'/clients/{cid}/edit',data={**payload,'csrf':ca,'version':va},follow_redirects=False); assert r.status_code==303
    r=b.post(f'/clients/{cid}/edit',data={**payload,'first_name':'Druhý','csrf':cb,'version':vb},follow_redirects=False); assert r.status_code==409
    with SessionLocal() as db: assert db.get(Client,cid).first_name=='První'


def test_reader_cannot_open_admin():
    c=TestClient(app); login(c,'reader@example.test')
    r=c.get('/admin/catalog'); assert r.status_code==403


def test_calendar_and_csv_export():
    c=TestClient(app); login(c)
    r=c.get('/daily-notes'); assert r.status_code==200 and 'Denní zápisy' in r.text
    r=c.get('/exports/reports.csv'); assert r.status_code==200 and 'text/csv' in r.headers['content-type']


def test_backup_and_restore(tmp_path):
    with SessionLocal() as db:
        run=perform_backup(db,str(tmp_path)); assert run.status=='ok'; assert Path(run.path).exists()
        rr=verify_restore(db,run.path); assert rr.status=='ok'


def synthetic_xlsx_bytes():
    from io import BytesIO
    wb=Workbook(); ws=wb.active; ws.title='Klienti'; ws.append(['KlientID','Kod','Jmeno','Prijmeni','DatumNarozeni','DatumVstupu','Kontakt','OmezenePrilezitosti','DokladUlozen','Stav','Poznamka']); ws.append(['CLI-IMP','IMP-1','Import','Test',date(1995,1,1),date(2026,10,1),'',None,False,'active',''])
    ws=wb.create_sheet('Vykony'); ws.append(['VykonID','KlientID','Datum','Minuty','PracovnikEmail','OblastID','CinnostID','FormaID','MistoID','MetodaID','JazykID','VysledekID','DalsiKrokID','TerminKroku','PokrokID','NavaznaSluzbaID','Popis']); ws.append(['VYK-IMP','CLI-IMP',date(2026,10,2),30,'import@example.test','OBL001','CIN001',None,None,None,None,None,None,None,None,None,'test'])
    ws=wb.create_sheet('Dokumenty'); ws.append(['DokumentID','KlientID','Nazev','Druh','Soubor','DatumVlozeni','PracovnikEmail','Poznamka'])
    ws=wb.create_sheet('DenniZapisy'); ws.append(['ZapisID','Datum','PracovnikEmail','ShrnutiDne'])
    ws=wb.create_sheet('Pracovnici'); ws.append(['Email','Jmeno','Role','Aktivni']); ws.append(['import@example.test','Import pracovník','Pracovnik',True])
    ws=wb.create_sheet('Ciselniky'); ws.append(['PolozkaID','Kategorie','NadrazenaID','Nazev','Poradi','Aktivni']); ws.append(['OBL001','Oblast',None,'Orientace a vstup do podpory',1,True]); ws.append(['CIN001','Cinnost','OBL001','První kontakt',1,True])
    bio=BytesIO(); wb.save(bio); return bio.getvalue()


def test_idempotent_import(tmp_path):
    with SessionLocal() as db:
        admin=get_worker(db); content=synthetic_xlsx_bytes(); batch,preview=save_import_upload(content,str(tmp_path),db,admin.id); assert not preview['errors']
        r1=apply_import(batch,db); assert r1['created'].get('Klienti')==1
        # Same file as a second batch: should not duplicate existing IDs.
        batch2,preview2=save_import_upload(content,str(tmp_path),db,admin.id); r2=apply_import(batch2,db)
        assert r2['unchanged'].get('Klienti')==1
        assert db.scalar(select(func.count()).select_from(Client).where(Client.client_id=='CLI-IMP'))==1
        assert db.scalar(select(func.count()).select_from(IndividualService).where(IndividualService.service_id=='VYK-IMP'))==1


def test_reader_privacy_and_aggregate_reporting():
    from app.indicators import indicator_summary
    with SessionLocal() as db:
        w=get_worker(db)
        c=make_client(db,'T-PRIV','Utajeny')
        db.add(IndicatorRecord(indicator_code='026300',client_id=c.id,record_date=date(2026,11,1),source_type='service',source_id='S-PRIV',eligible=True))
        db.commit()
        assert any(x['code']=='026300' and x['value'] >= 1 for x in indicator_summary(db,date(2026,11,1),date(2026,11,1)))
    tc=TestClient(app); login(tc,'reader@example.test')
    assert tc.get('/clients').status_code==403
    assert tc.get('/activity-report').status_code==403
    reports=tc.get('/reports')
    assert reports.status_code==200
    assert '/activity-report' not in reports.text
    detail=tc.get('/reports/026300?start=2026-11-01&end=2026-11-01')
    assert detail.status_code==200
    assert 'T-PRIV' in detail.text
    assert 'Utajeny Klient' not in detail.text
    assert '/clients/' not in detail.text


def test_attendance_sheet_indicator_is_unique_and_traceable():
    day=date(2026,11,10)
    with SessionLocal() as db:
        w=get_worker(db)
        clients=[make_client(db,f'T-ATT{i}',f'Att{i}') for i in range(1,4)]
        e=Event(event_id='E-ATT-1',title='Občanská orientace',event_date=day,ka='KA2',subactivity='2.3',event_type='civic_orientation',attendance_sheet_ref='PL-2026-11-10',created_by_id=w.id)
        db.add(e); db.flush()
        for c in clients:
            db.add(EventParticipant(event_id=e.id,client_id=c.id,participant_kind='project',attendance_status='present',eligible_project=True,transcribed_from_sheet=True))
        db.add(EventParticipant(event_id=e.id,client_id=None,display_name='host bez projektové evidence',participant_kind='guest',attendance_status='present',eligible_project=False))
        sync_indicator_records(db,'event',e.event_id,e.event_date,[c.id for c in clients],['026200'],True,e.attendance_sheet_ref)
        # Druhá účast téže osoby nesmí zvýšit počet unikátních osob.
        e2=Event(event_id='E-ATT-2',title='Občanská orientace II',event_date=day,ka='KA2',subactivity='2.3',event_type='civic_orientation',attendance_sheet_ref='PL-2026-11-10-B',created_by_id=w.id)
        db.add(e2); db.flush()
        db.add(EventParticipant(event_id=e2.id,client_id=clients[0].id,participant_kind='project',attendance_status='present',eligible_project=True,transcribed_from_sheet=True))
        sync_indicator_records(db,'event',e2.event_id,e2.event_date,[clients[0].id],['026200'],True,e2.attendance_sheet_ref)
        db.commit()
        ids=indicator_clients(db,'026200',day,day)
        assert set(ids)=={c.id for c in clients}
    tc=TestClient(app); login(tc)
    r=tc.get('/reports/026200?start=2026-11-10&end=2026-11-10')
    assert r.status_code==200
    assert 'PL-2026-11-10' in r.text and 'PL-2026-11-10-B' in r.text
    assert 'host bez projektové evidence' not in r.text


def test_ineligible_service_cannot_leave_specific_indicator():
    day=date(2026,11,11)
    with SessionLocal() as db:
        w=get_worker(db); c=make_client(db,'T-INEL','Nezpusobily')
        s=IndividualService(service_id='S-INEL',client_id=c.id,service_date=day,minutes=25,worker_id=w.id,ka='KA3',subactivity='3.1',eligible_project=True)
        db.add(s); db.flush()
        sync_indicator_records(db,'service',s.service_id,day,[c.id],['026300'],False,None)
        db.commit()
        assert c.id in indicator_clients(db,'026300',day,day)
        s.eligible_project=False
        sync_indicator_records(db,'service',s.service_id,day,[c.id],[],False,None)
        db.commit()
        assert c.id not in indicator_clients(db,'026300',day,day)


def test_multiple_daily_notes_and_edit_history():
    tc=TestClient(app); login(tc)
    token=csrf(tc,'/daily-notes/new?date=2026-11-12')
    for text in ('První zápis dne','Druhý zápis dne'):
        r=tc.post('/daily-notes/new',data={'csrf':token,'note_date':'2026-11-12','summary':text},follow_redirects=False)
        assert r.status_code==303
    with SessionLocal() as db:
        notes=db.scalars(select(DailyNote).where(DailyNote.note_date==date(2026,11,12)).order_by(DailyNote.id)).all()
        assert len(notes)==2
        nid=notes[0].id
    page=tc.get(f'/daily-notes/{nid}/edit')
    token2=re.search(r'name="csrf" value="([^"]+)"',page.text).group(1)
    ver=re.search(r'name="version" value="(\d+)"',page.text).group(1)
    r=tc.post(f'/daily-notes/{nid}/edit',data={'csrf':token2,'version':ver,'note_date':'2026-11-12','summary':'První zápis – opravený'},follow_redirects=False)
    assert r.status_code==303
    with SessionLocal() as db:
        n=db.get(DailyNote,nid)
        assert n.summary=='První zápis – opravený'
        hist=db.scalars(select(AuditLog).where(AuditLog.entity_type=='daily_note',AuditLog.entity_id==n.note_id,AuditLog.action=='update')).all()
        assert hist


def test_activity_xlsx_contains_filter_metadata():
    from io import BytesIO
    from openpyxl import load_workbook
    tc=TestClient(app); login(tc)
    r=tc.get('/exports/activity.xlsx?start=2026-11-01&end=2026-11-30&ka=KA3')
    assert r.status_code==200
    wb=load_workbook(BytesIO(r.content),read_only=True,data_only=True)
    assert 'Filtry' in wb.sheetnames
    vals={str(row[0]): row[1] for row in wb['Filtry'].iter_rows(min_row=2,values_only=True) if row[0]}
    assert vals['start']=='2026-11-01'
    assert vals['end']=='2026-11-30'
    assert vals['ka']=='KA3'


def test_project_outputs_web_visits_and_anonymous_feedback_are_separate():
    from app.indicators import indicator_summary
    day=date(2026,11,15)
    with SessionLocal() as db:
        w=get_worker(db)
        for i in range(5): db.add(ProjectOutput(output_date=day,title=f'Dokument {i+1}',output_type='document',created_by_id=w.id))
        db.add(WebAnalyticsAggregate(period_start=date(2026,11,1),period_end=date(2026,11,30),source='web',visits=123,created_by_id=w.id))
        db.add(AnonymousFeedbackAggregate(period_start=date(2026,11,1),period_end=date(2026,11,30),denominator=10,positive_count=7,methodology='Anonymní dotazník – test',created_by_id=w.id))
        db.commit()
        summary={x['code']:x['value'] for x in indicator_summary(db,date(2026,11,1),date(2026,11,30))}
        assert summary['028309']>=5
        assert summary['028313']>=123
        # Bez schválené metodiky se anonymní 70% cíl nesmí automaticky vydávat za indikátor 026510.
        assert summary['026510']==0


def test_catalog_area_with_children_cannot_be_physically_deleted():
    tc=TestClient(app); login(tc)
    with SessionLocal() as db:
        area=db.scalar(select(CatalogItem).where(CatalogItem.item_id=='OBL001'))
        assert area and db.scalar(select(func.count()).select_from(CatalogItem).where(CatalogItem.parent_item_id=='OBL001'))>0
        pk=area.id
    token=csrf(tc,'/admin/catalog')
    r=tc.post(f'/admin/catalog/{pk}/delete',data={'csrf':token},follow_redirects=False)
    assert r.status_code==409


def test_import_warns_about_missing_ka_and_uses_legacy_marker(tmp_path):
    with SessionLocal() as db:
        admin=get_worker(db); batch,preview=save_import_upload(synthetic_xlsx_bytes(),str(tmp_path),db,admin.id)
        assert preview['warnings'] and 'KA' in preview['warnings'][0]
        apply_import(batch,db)
        svc=db.scalar(select(IndividualService).where(IndividualService.service_id=='VYK-IMP'))
        assert svc.ka=='legacy'
