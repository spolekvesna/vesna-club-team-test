from __future__ import annotations
import hashlib, json, shutil
from datetime import datetime, date, timezone
from pathlib import Path
from uuid import uuid4
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session
from .models import ImportBatch, CatalogItem, CatalogCategory, Client, Worker, IndividualService, Document, DailyNote
from .catalog_seed import CATEGORY_NAMES
from .security import ROLE_ADMIN, ROLE_WORKER, ROLE_READER
from .utils import parse_date, uid

EXPECTED = {
    'Klienti':['KlientID','Kod','Jmeno','Prijmeni','DatumNarozeni','DatumVstupu','Kontakt','OmezenePrilezitosti','DokladUlozen','Stav','Poznamka'],
    'Vykony':['VykonID','KlientID','Datum','Minuty','PracovnikEmail','OblastID','CinnostID','FormaID','MistoID','MetodaID','JazykID','VysledekID','DalsiKrokID','TerminKroku','PokrokID','NavaznaSluzbaID','Popis'],
    'Dokumenty':['DokumentID','KlientID','Nazev','Druh','Soubor','DatumVlozeni','PracovnikEmail','Poznamka'],
    'DenniZapisy':['ZapisID','Datum','PracovnikEmail','ShrnutiDne'],
    'Pracovnici':['Email','Jmeno','Role','Aktivni'],
    'Ciselniky':['PolozkaID','Kategorie','NadrazenaID','Nazev','Poradi','Aktivni'],
}

def _norm_date(v):
    if not v: return None
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    return date.fromisoformat(str(v)[:10])

def preview_xlsx(path: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    errors=[]; warnings=[]; counts={}; headers={}
    for sheet, expected in EXPECTED.items():
        if sheet not in wb.sheetnames:
            errors.append(f"Chybí list {sheet}")
            continue
        ws=wb[sheet]
        actual=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
        headers[sheet]=actual
        if actual[:len(expected)] != expected:
            errors.append(f"List {sheet}: neočekávané hlavičky")
        counts[sheet]=max(ws.max_row-1,0)
        if sheet == 'Vykony' and counts[sheet] > 0:
            warnings.append('List Vykony neobsahuje KA ani podaktivitu. Importované výkony budou označeny jako legacy a musí se před vykazováním zkontrolovat/doplnit.')
    catalog_counts={}
    if 'Ciselniky' in wb.sheetnames:
        ws=wb['Ciselniky']
        for row in ws.iter_rows(min_row=2,values_only=True):
            if not row[0]: continue
            catalog_counts[row[1]]=catalog_counts.get(row[1],0)+1
    return {"counts":counts,"headers":headers,"catalog_counts":catalog_counts,"errors":errors,"warnings":warnings}

def save_import_upload(content: bytes, imports_dir: str, db: Session, user_id: int):
    Path(imports_dir).mkdir(parents=True, exist_ok=True)
    token=uuid4().hex
    path=Path(imports_dir)/f"{token}.xlsx"
    path.write_bytes(content)
    sha=hashlib.sha256(content).hexdigest()
    preview=preview_xlsx(str(path))
    batch=ImportBatch(token=token,file_path=str(path),sha256=sha,status='preview',preview_json=json.dumps(preview,ensure_ascii=False),created_by_id=user_id)
    db.add(batch); db.commit(); db.refresh(batch)
    return batch, preview

def _bool(v): return str(v).lower() in {'true','1','yes','ano'} if not isinstance(v,bool) else v

def apply_import(batch: ImportBatch, db: Session):
    wb=load_workbook(batch.file_path,data_only=True)
    stats={"created":{},"updated":{},"unchanged":{},"errors":[]}
    def bump(kind,key): stats[kind][key]=stats[kind].get(key,0)+1
    # catalog categories/items
    for i,key in enumerate(CATEGORY_NAMES,start=1):
        cat=db.scalar(select(CatalogCategory).where(CatalogCategory.key==key))
        if not cat:
            db.add(CatalogCategory(key=key,name=CATEGORY_NAMES[key],sort_order=i)); bump('created','Ciselniky-kategorie')
    db.flush()
    ws=wb['Ciselniky']; heads=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]: continue
        d=dict(zip(heads,row)); obj=db.scalar(select(CatalogItem).where(CatalogItem.item_id==d['PolozkaID']))
        vals=dict(category_key=d['Kategorie'],parent_item_id=d['NadrazenaID'] or None,name=d['Nazev'],sort_order=int(d['Poradi'] or 0),active=_bool(d['Aktivni']))
        if not obj:
            db.add(CatalogItem(item_id=d['PolozkaID'],visible=True,**vals)); bump('created','Ciselniky')
        elif any(getattr(obj,k)!=v for k,v in vals.items()):
            for k,v in vals.items(): setattr(obj,k,v)
            obj.version+=1; bump('updated','Ciselniky')
        else: bump('unchanged','Ciselniky')
    # workers. Production import may carry real emails; never generated into seed/test data.
    ws=wb['Pracovnici']; heads=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]: continue
        d=dict(zip(heads,row)); role_raw=str(d['Role'] or '').lower()
        if 'admin' in role_raw or 'správ' in role_raw: role=ROLE_ADMIN
        elif 'reader' in role_raw or 'čten' in role_raw or 'report' in role_raw: role=ROLE_READER
        else: role=ROLE_WORKER
        obj=db.scalar(select(Worker).where(Worker.email==d['Email']))
        vals={'name':d['Jmeno'] or d['Email'],'role':role,'active':_bool(d['Aktivni'])}
        if not obj:
            db.add(Worker(email=d['Email'],**vals)); bump('created','Pracovnici')
        elif any(getattr(obj,k)!=v for k,v in vals.items()):
            for k,v in vals.items(): setattr(obj,k,v)
            obj.version+=1; bump('updated','Pracovnici')
        else: bump('unchanged','Pracovnici')
    db.flush()
    # clients
    ws=wb['Klienti']; heads=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]: continue
        d=dict(zip(heads,row)); obj=db.scalar(select(Client).where(Client.client_id==d['KlientID']))
        if not obj:
            
            raw_doc=d['DokladUlozen']; doc_status='unknown' if raw_doc in (None,'') else ('yes' if _bool(raw_doc) else 'no')
            db.add(Client(client_id=d['KlientID'],code=d['Kod'],first_name=d['Jmeno'],last_name=d['Prijmeni'],birth_date=_norm_date(d['DatumNarozeni']),entry_date=_norm_date(d['DatumVstupu']),contact=d['Kontakt'],fewer_opportunities=d['OmezenePrilezitosti'] if isinstance(d['OmezenePrilezitosti'],bool) else None,document_status=doc_status,status=d['Stav'] or 'active',notes=d['Poznamka'])); bump('created','Klienti')
        else: bump('unchanged','Klienti')
    db.flush()
    # performance records
    ws=wb['Vykony']; heads=[c.value for c in ws[1]]
    clients={x.client_id:x for x in db.scalars(select(Client)).all()}; workers={x.email:x for x in db.scalars(select(Worker)).all()}
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]: continue
        d=dict(zip(heads,row)); obj=db.scalar(select(IndividualService).where(IndividualService.service_id==d['VykonID']))
        if obj: bump('unchanged','Vykony'); continue
        c=clients.get(d['KlientID']); w=workers.get(d['PracovnikEmail'])
        if not c or not w:
            stats['errors'].append(f"Výkon {d['VykonID']}: chybí klient nebo pracovník"); continue
        vals={'area_id':d['OblastID'],'activity_id':d['CinnostID'],'form_id':d['FormaID'],'place_id':d['MistoID'],'method_id':d['MetodaID'],'language_id':d['JazykID'],'outcome_id':d['VysledekID'],'next_step_id':d['DalsiKrokID'],'progress_id':d['PokrokID'],'referral_id':d['NavaznaSluzbaID']}
        ids=[v for v in vals.values() if v]; name_map={x.item_id:x.name for x in db.scalars(select(CatalogItem).where(CatalogItem.item_id.in_(ids))).all()} if ids else {}
        snap=json.dumps({k:{'id':v,'name':name_map.get(v)} for k,v in vals.items() if v},ensure_ascii=False)
        db.add(IndividualService(service_id=d['VykonID'],client_id=c.id,service_date=_norm_date(d['Datum']),minutes=int(d['Minuty'] or 0),worker_id=w.id,ka='legacy',area_id=d['OblastID'],activity_id=d['CinnostID'],form_id=d['FormaID'],place_id=d['MistoID'],method_id=d['MetodaID'],language_id=d['JazykID'],outcome_id=d['VysledekID'],next_step_id=d['DalsiKrokID'],follow_up_date=_norm_date(d['TerminKroku']),progress_id=d['PokrokID'],referral_id=d['NavaznaSluzbaID'],description=d['Popis'],catalog_snapshot_json=snap,eligible_project=True)); bump('created','Vykony')
    db.flush()
    # Documents and daily notes
    ws=wb['Dokumenty']; heads=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]: continue
        d=dict(zip(heads,row));
        if db.scalar(select(Document).where(Document.document_id==d['DokumentID'])): bump('unchanged','Dokumenty'); continue
        c=clients.get(d['KlientID']); w=workers.get(d['PracovnikEmail'])
        if c and w:
            db.add(Document(document_id=d['DokumentID'],client_id=c.id,name=d['Nazev'],kind=d['Druh'],drive_ref=d['Soubor'],inserted_date=_norm_date(d['DatumVlozeni']) or date.today(),worker_id=w.id,notes=d['Poznamka'])); bump('created','Dokumenty')
        else: stats['errors'].append(f"Dokument {d['DokumentID']}: chybí klient nebo pracovník")
    ws=wb['DenniZapisy']; heads=[c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2,values_only=True):
        if not row[0]: continue
        d=dict(zip(heads,row));
        if db.scalar(select(DailyNote).where(DailyNote.note_id==d['ZapisID'])): bump('unchanged','DenniZapisy'); continue
        w=workers.get(d['PracovnikEmail'])
        if w:
            db.add(DailyNote(note_id=d['ZapisID'],note_date=_norm_date(d['Datum']),worker_id=w.id,summary=d['ShrnutiDne'] or '')); bump('created','DenniZapisy')
        else: stats['errors'].append(f"Denní zápis {d['ZapisID']}: chybí pracovník")
    batch.status='applied'; batch.applied_at=datetime.now(timezone.utc).replace(tzinfo=None)
    db.commit()
    return stats
